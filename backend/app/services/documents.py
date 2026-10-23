"""Render a MOM as a Word file or a PDF.

Only the languages the reviewer kept are written out. The others stay in the stored
document untouched, so exporting never changes what exists.

Section labels follow the language of the section they introduce, because a client
reading only the Chinese part should not have to read English headings.
"""

from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.document import Document as DocxDocument
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

from app.schemas.meeting import Language
from app.schemas.mom import Audience, MomAction, MomContent, MomDocument, MomVersion

LANGUAGE_ORDER = [Language.en, Language.zh, Language.id]

LABELS: dict[Language, dict[str, str]] = {
    Language.en: {
        "name": "English",
        "overview": "Overview",
        "topics": "Topics",
        "decisions": "Decisions",
        "actions": "Action items",
        "questions": "Open questions",
        "notes": "Internal notes",
        "owner": "Owner",
        "deadline": "Due",
    },
    Language.zh: {
        "name": "中文",
        "overview": "会议概述",
        "topics": "讨论议题",
        "decisions": "决定",
        "actions": "行动项",
        "questions": "待定问题",
        "notes": "内部备注",
        "owner": "负责人",
        "deadline": "截止日期",
    },
    Language.id: {
        "name": "Bahasa Indonesia",
        "overview": "Ringkasan",
        "topics": "Topik",
        "decisions": "Keputusan",
        "actions": "Tindakan",
        "questions": "Pertanyaan terbuka",
        "notes": "Catatan internal",
        "owner": "Penanggung jawab",
        "deadline": "Tenggat",
    },
}

# STSong-Light ships with reportlab and covers Chinese, so no font file is needed. Its
# Latin metrics are wrong, though, so Latin text is set in Helvetica and only Chinese
# runs switch font. Setting the whole paragraph in the CID font overlaps Latin words.
PDF_CJK_FONT = "STSong-Light"
PDF_LATIN_FONT = "Helvetica"

CJK_RANGES = (
    (0x2E80, 0x9FFF),
    (0x3000, 0x303F),
    (0x3400, 0x4DBF),
    (0xFF00, 0xFFEF),
)


def _is_cjk(character: str) -> bool:
    code = ord(character)
    return any(low <= code <= high for low, high in CJK_RANGES)


def _markup(text: str) -> str:
    """Escape for reportlab, wrapping only the Chinese runs in the CJK font."""
    runs: list[tuple[bool, list[str]]] = []
    for character in text:
        cjk = _is_cjk(character)
        if runs and runs[-1][0] is cjk:
            runs[-1][1].append(character)
        else:
            runs.append((cjk, [character]))

    parts = []
    for cjk, characters in runs:
        chunk = escape("".join(characters))
        parts.append(f'<font name="{PDF_CJK_FONT}">{chunk}</font>' if cjk else chunk)
    return "".join(parts)


def selected_versions(document: MomDocument) -> list[MomVersion]:
    """The versions to print, in a fixed order regardless of storage order."""
    by_language = {version.language: version for version in document.versions}
    return [
        by_language[language]
        for language in LANGUAGE_ORDER
        if language in set(document.visible_languages)
    ]


def _action_line(action: MomAction, labels: dict[str, str]) -> str:
    parts = [action.task]
    if action.owner:
        parts.append(f"{labels['owner']}: {action.owner}")
    if action.deadline:
        parts.append(f"{labels['deadline']}: {action.deadline}")
    return " · ".join(parts)


def _sections(
    content: MomContent, labels: dict[str, str], audience: Audience
) -> list[tuple[str, list[str]]]:
    sections = [
        (
            labels["topics"],
            [f"{topic.heading}: {topic.discussion}" for topic in content.topics],
        ),
        (labels["decisions"], list(content.decisions)),
        (
            labels["actions"],
            [_action_line(action, labels) for action in content.action_items],
        ),
        (labels["questions"], list(content.open_questions)),
    ]
    if audience is Audience.internal:
        sections.append((labels["notes"], list(content.internal_notes)))
    return [(heading, items) for heading, items in sections if items]


def _docx_section(doc: DocxDocument, heading: str, items: list[str]) -> None:
    doc.add_heading(heading, level=2)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def to_docx(document: MomDocument) -> bytes:
    doc = Document()
    for index, version in enumerate(selected_versions(document)):
        if index:
            doc.add_page_break()
        labels = LABELS[version.language]
        doc.add_heading(version.content.title, level=1)
        doc.add_paragraph(labels["name"]).italic = True
        if version.content.overview:
            doc.add_paragraph(version.content.overview)
        for heading, items in _sections(
            version.content, labels, document.audience
        ):
            _docx_section(doc, heading, items)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _styles() -> dict[str, ParagraphStyle]:
    return {
        "title": ParagraphStyle(
            "MomTitle", fontName=PDF_LATIN_FONT, fontSize=16, leading=22
        ),
        "label": ParagraphStyle(
            "MomLabel",
            fontName=PDF_LATIN_FONT,
            fontSize=9,
            leading=12,
            textColor="#555555",
        ),
        "heading": ParagraphStyle(
            "MomHeading",
            fontName=PDF_LATIN_FONT,
            fontSize=12,
            leading=16,
            spaceBefore=10,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "MomBody", fontName=PDF_LATIN_FONT, fontSize=10, leading=14
        ),
        "bullet": ParagraphStyle(
            "MomBullet",
            fontName=PDF_LATIN_FONT,
            fontSize=10,
            leading=14,
            leftIndent=12,
            bulletIndent=0,
        ),
    }


def to_pdf(document: MomDocument) -> bytes:
    pdfmetrics.registerFont(UnicodeCIDFont(PDF_CJK_FONT))
    styles = _styles()
    buffer = BytesIO()
    frame = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title="Minutes of meeting",
    )

    story: list[object] = []
    for index, version in enumerate(selected_versions(document)):
        if index:
            story.append(PageBreak())
        labels = LABELS[version.language]
        story.append(Paragraph(_markup(version.content.title), styles["title"]))
        story.append(Paragraph(_markup(labels["name"]), styles["label"]))
        story.append(Spacer(1, 6))
        if version.content.overview:
            story.append(Paragraph(_markup(version.content.overview), styles["body"]))
        for heading, items in _sections(
            version.content, labels, document.audience
        ):
            story.append(Paragraph(_markup(heading), styles["heading"]))
            for item in items:
                story.append(
                    Paragraph(_markup(item), styles["bullet"], bulletText="•")
                )

    frame.build(story)
    return buffer.getvalue()
