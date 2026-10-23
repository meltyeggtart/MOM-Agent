from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.schemas.meeting import Language
from app.schemas.mom import Audience, MomContent, MomDocument, MomVersion
from app.services.documents import LABELS, selected_versions, to_docx, to_pdf

TITLES = {
    Language.en: "Q3 kickoff",
    Language.zh: "第三季度启动会",
    Language.id: "Kickoff Q3",
}


def content(language: Language, notes: list[str] | None = None) -> MomContent:
    return MomContent(
        title=TITLES[language],
        overview="The team kept the budget.",
        topics=[{"heading": "Q3 campaign", "discussion": "The launch moves."}],
        decisions=["Keep the budget"],
        action_items=[
            {"task": "Send the plan", "owner": "Ana", "deadline": "15 November 2026"}
        ],
        open_questions=["Who owns the review?"],
        internal_notes=notes if notes is not None else ["Margin is thin."],
    )


def document(
    visible: list[Language] | None = None, audience: Audience = Audience.internal
) -> MomDocument:
    notes = ["Margin is thin."] if audience is Audience.internal else []
    return MomDocument(
        source_language=Language.en,
        audience=audience,
        visible_languages=visible or list(Language),
        versions=[
            MomVersion(language=language, content=content(language, notes))
            for language in Language
        ],
    )


def docx_text(data: bytes) -> str:
    return "\n".join(paragraph.text for paragraph in Document(BytesIO(data)).paragraphs)


def pdf_text(data: bytes) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)


def test_labels_exist_for_every_language() -> None:
    assert set(LABELS) == set(Language)


def test_only_visible_languages_are_selected() -> None:
    selected = selected_versions(document([Language.id, Language.en]))

    assert [version.language for version in selected] == [Language.en, Language.id]


def test_docx_carries_every_visible_language() -> None:
    text = docx_text(to_docx(document()))

    for title in TITLES.values():
        assert title in text


def test_docx_omits_a_hidden_language() -> None:
    text = docx_text(to_docx(document([Language.en, Language.id])))

    assert TITLES[Language.en] in text
    assert TITLES[Language.id] in text
    assert TITLES[Language.zh] not in text


def test_docx_uses_the_language_for_section_labels() -> None:
    text = docx_text(to_docx(document()))

    assert LABELS[Language.zh]["decisions"] in text
    assert LABELS[Language.id]["decisions"] in text


def test_external_export_has_no_internal_notes() -> None:
    external = document(audience=Audience.external)

    docx = docx_text(to_docx(external))
    pdf = pdf_text(to_pdf(external))

    for text in (docx, pdf):
        assert "Margin is thin." not in text
        assert LABELS[Language.en]["notes"] not in text


def test_internal_export_keeps_internal_notes() -> None:
    assert "Margin is thin." in docx_text(to_docx(document()))


def test_pdf_has_one_page_per_language() -> None:
    reader = PdfReader(BytesIO(to_pdf(document())))

    assert len(reader.pages) == 3


def test_pdf_keeps_chinese_readable() -> None:
    text = pdf_text(to_pdf(document()))

    assert TITLES[Language.zh] in text
    assert LABELS[Language.zh]["decisions"] in text
    assert TITLES[Language.en] in text


def test_pdf_drops_a_hidden_language() -> None:
    text = pdf_text(to_pdf(document([Language.en, Language.id])))

    assert TITLES[Language.zh] not in text
    assert len(PdfReader(BytesIO(to_pdf(document([Language.en])))).pages) == 1


def create_meeting(client: TestClient) -> str:
    response = client.post(
        "/meetings",
        json={
            "title": "Q3 kickoff",
            "meeting_at": "2026-10-20T09:00:00Z",
            "languages": ["en"],
        },
    )
    return response.json()["id"]


def store(client: TestClient, meeting_id: str) -> None:
    response = client.post(
        f"/meetings/{meeting_id}/moms/internal/versions",
        json=document().model_dump(mode="json"),
    )
    assert response.status_code == 201


def test_endpoint_returns_a_word_file(client: TestClient) -> None:
    meeting_id = create_meeting(client)
    store(client, meeting_id)

    response = client.get(f"/meetings/{meeting_id}/moms/internal/export.docx")

    assert response.status_code == 200
    assert response.content.startswith(b"PK")
    assert "wordprocessingml" in response.headers["content-type"]
    assert f'mom-{meeting_id}-internal.docx' in response.headers["content-disposition"]


def test_endpoint_returns_a_pdf(client: TestClient) -> None:
    meeting_id = create_meeting(client)
    store(client, meeting_id)

    response = client.get(f"/meetings/{meeting_id}/moms/internal/export.pdf")

    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert response.headers["content-type"] == "application/pdf"
    assert TITLES[Language.zh] in pdf_text(response.content)


def test_export_of_an_unknown_meeting_is_404(client: TestClient) -> None:
    missing = "00000000-0000-0000-0000-000000000000"

    assert (
        client.get(f"/meetings/{missing}/moms/internal/export.docx").status_code == 404
    )
    assert client.get(f"/meetings/{missing}/moms/internal/export.pdf").status_code == 404
